import os
os.environ['DB_ENV'] = 'developmentAWS'

import pytest
import requests
from io import BytesIO
try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None
from app import create_app
from config import Config

# Helpers (misma filosofía que en tests existentes)
@pytest.fixture
def client():
    Config.set_environment_variables()
    app = create_app()
    app.config['TESTING'] = True
    with app.app_context():
        yield app.test_client()

BASE_URL = "http://127.0.0.1:5000"


def login_and_get_tokens(email, password):
    login_response = requests.post(f"{BASE_URL}/auth/login", json={'email': email, 'password': password})
    assert login_response.status_code == 200
    data = login_response.json()
    tokens = data.get('tokens') or {}
    return tokens.get('access_token'), tokens.get('refresh_token')


def list_users(access_token, params=None):
    headers = {'Authorization': f'Bearer {access_token}'} if access_token else {}
    qs = ''
    if params:
        from urllib.parse import urlencode
        qs = '?' + urlencode(params)
    return requests.get(f"{BASE_URL}/usuarios/{qs}", headers=headers)


def export_users(access_token, fmt='csv', params=None):
    headers = {'Authorization': f'Bearer {access_token}'} if access_token else {}
    qs = f'?format={fmt}'
    if params:
        from urllib.parse import urlencode
        qs += '&' + urlencode(params)
    return requests.get(f"{BASE_URL}/usuarios/export{qs}", headers=headers, stream=True)


def logout_with_refresh(refresh_token):
    return requests.post(f"{BASE_URL}/auth/logout", headers={'Authorization': f'Bearer {refresh_token}'})


# ------------------ Tests para Admin Plataforma ------------------

@pytest.mark.meta(title='Admin plataforma - listar todos los usuarios', desc='Login admin_plataforma y listar todos los usuarios')
def test_admin_plataforma_list_all_users(client):
    print("\nPrueba: admin_plataforma inicia sesión y lista todos los usuarios de la plataforma.")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = list_users(access)
    assert rv.status_code == 200
    resp = rv.json()
    assert isinstance(resp, dict)
    # Comprobar estructura del envelope de paginación y valores básicos
    assert 'items' in resp and 'page' in resp and 'size' in resp and 'returned' in resp and 'has_more' in resp
    # Comprobar tipos y coherencia de campos
    assert isinstance(resp['items'], list)
    assert isinstance(resp['page'], int)
    assert isinstance(resp['size'], int)
    assert isinstance(resp['returned'], int)
    assert isinstance(resp['has_more'], bool)
    # Si size > 0, returned no debe exceder size
    if resp['size'] > 0:
        assert resp['returned'] <= resp['size']
    # next_page/prev_page pueden no estar presentes, pero si lo están deben ser ints o None
    if 'next_page' in resp and resp['next_page'] is not None:
        assert isinstance(resp['next_page'], int)
    if 'prev_page' in resp and resp['prev_page'] is not None:
        assert isinstance(resp['prev_page'], int)


@pytest.mark.meta(title='Admin plataforma - export a Excel (XLSX)', desc='Exportar todos los usuarios a XLSX')
def test_admin_plataforma_export_xlsx(client):
    print("\nPrueba: admin_plataforma exporta todos los usuarios a Excel (XLSX).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = export_users(access, fmt='xlsx')
    # Para archivos binarios esperamos 200 y content-type apropiado
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in ctype
    # Validar contenido del XLSX en memoria si openpyxl está disponible
    if load_workbook:
        wb = load_workbook(filename=BytesIO(rv.content), read_only=True)
        # Esperamos al menos una hoja llamada 'usuarios' o la primera hoja con datos
        ws = None
        if 'usuarios' in wb.sheetnames:
            ws = wb['usuarios']
        else:
            ws = wb[wb.sheetnames[0]]
        # Leer la primera fila como cabecera
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 1
        header = [str(c).lower() if c is not None else '' for c in rows[0]]
        # Cabeceras esperadas
        expected_cols = {'id', 'nombre', 'email', 'rol', 'academia_id', 'estado', 'fecha_alta'}
        assert expected_cols.issubset(set(header)), f"Cabeceras XLSX inesperadas: {header}"


@pytest.mark.meta(title='Admin plataforma - export a CSV', desc='Exportar todos los usuarios a CSV')
def test_admin_plataforma_export_csv(client):
    print("\nPrueba: admin_plataforma exporta todos los usuarios a CSV.")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = export_users(access, fmt='csv')
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'text/csv' in ctype
    # Comprobar cabecera CSV
    content = rv.content.decode('utf-8')
    assert 'id' in content.splitlines()[0].lower()
    # Comprobar que el número de líneas coincide con el contado por la API (si with_total no usado no se puede asegurar)
    lines = content.splitlines()
    assert len(lines) >= 1


@pytest.mark.meta(title='Admin plataforma - busqueda por academia_id=303', desc='Listar usuarios filtrando por academia_id=303')
def test_admin_plataforma_list_academia_303(client):
    print("\nPrueba: admin_plataforma lista usuarios filtrando por academia_id=303.")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = list_users(access, params={'academia_id': 303})
    # Asumimos que la ruta responde 200 aunque la lista pueda estar vacía
    assert rv.status_code == 200
    resp = rv.json()
    assert 'items' in resp


@pytest.mark.meta(title='Admin plataforma - export academia 303 a CSV', desc='Exportar usuarios de academia 303 a CSV')
def test_admin_plataforma_export_academia_303_csv(client):
    print("\nPrueba: admin_plataforma exporta usuarios de academia_id=303 a CSV.")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = export_users(access, fmt='csv', params={'academia_id': 303})
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'text/csv' in ctype


@pytest.mark.meta(title='Admin plataforma - export academia 303 a XLSX', desc='Exportar usuarios de academia 303 a XLSX')
def test_admin_plataforma_export_academia_303_xlsx(client):
    print("\nPrueba: admin_plataforma exporta usuarios de academia_id=303 a Excel (XLSX).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = export_users(access, fmt='xlsx', params={'academia_id': 303})
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in ctype
    if load_workbook:
        wb = load_workbook(filename=BytesIO(rv.content), read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        # Al menos la cabecera
        assert len(rows) >= 1
        header = [str(c).lower() if c is not None else '' for c in rows[0]]
        expected_cols = {'id', 'nombre', 'email', 'rol', 'academia_id', 'estado', 'fecha_alta'}
        assert expected_cols.issubset(set(header))


@pytest.mark.meta(title='Admin plataforma - logout', desc='Logout del admin_plataforma')
def test_admin_plataforma_logout(client):
    print("\nPrueba: admin_plataforma realiza logout (revoca refresh token).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = logout_with_refresh(refresh)
    assert rv.status_code in (200, 204)


# Test de paginación básica para Admin Plataforma
@pytest.mark.meta(title='Admin plataforma - paginación básica', desc='Comprobar paginación básica con page=1 y size=10')
def test_admin_plataforma_pagination_basic(client):
    print("\nPrueba: admin_plataforma verifica paginación básica (page=1, size=10).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = list_users(access, params={'page': 1, 'size': 10})
    assert rv.status_code == 200
    resp = rv.json()
    assert len(resp['items']) <= 10
    assert resp['page'] == 1
    assert resp['size'] == 10
    assert isinstance(resp['has_more'], bool)


# Test de paginación avanzada para Admin Plataforma
@pytest.mark.meta(title='Admin plataforma - paginación avanzada', desc='Comprobar paginación avanzada con page=2 y size=5')
def test_admin_plataforma_pagination_advanced(client):
    print("\nPrueba: admin_plataforma verifica paginación avanzada (page=2, size=5).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = list_users(access, params={'page': 2, 'size': 5})
    assert rv.status_code == 200
    resp = rv.json()
    assert len(resp['items']) <= 5
    assert resp['page'] == 2
    assert resp['size'] == 5
    assert isinstance(resp['has_more'], bool)


# Test de paginación con filtros para Admin Plataforma
@pytest.mark.meta(title='Admin plataforma - paginación con filtros', desc='Comprobar paginación con filtros y academia_id=303')
def test_admin_plataforma_pagination_with_filters(client):
    print("\nPrueba: admin_plataforma verifica paginación con filtros (academia_id=303).")
    access, refresh = login_and_get_tokens('admin_plataforma@academia.com', 'password_admin_plataforma')
    rv = list_users(access, params={'page': 1, 'size': 5, 'academia_id': 303})
    assert rv.status_code == 200
    resp = rv.json()
    assert len(resp['items']) <= 5
    for user in resp['items']:
        assert user['academia_id'] == 303


# ------------------ Tests para Admin Academia (no vinculado a academy 303) ------------------

@pytest.mark.meta(title='Admin_academia - listar sus usuarios', desc='Login admin_academia y listar usuarios; solo los de su academia')
def test_admin_academia_list_own_users_only(client):
    print("\nPrueba: admin_academia inicia sesión y lista solo los usuarios de su academia.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = list_users(access)
    assert rv.status_code == 200
    resp = rv.json()
    users = resp.get('items', [])
    # Obtener academia del propio usuario
    me = requests.get(f"{BASE_URL}/usuarios/me", headers={'Authorization': f'Bearer {access}'}).json()
    my_academia = me.get('academia_id')
    for u in users:
        assert u.get('academia_id') == my_academia


@pytest.mark.meta(title='Admin_academia - export a XLSX (propia academia)', desc='Exportar usuarios de su academia a XLSX')
def test_admin_academia_export_own_xlsx(client):
    print("\nPrueba: admin_academia exporta los usuarios de su propia academia a XLSX.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = export_users(access, fmt='xlsx')
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in ctype


@pytest.mark.meta(title='Admin_academia - export a CSV (propia academia)', desc='Exportar usuarios de su academia a CSV')
def test_admin_academia_export_own_csv(client):
    print("\nPrueba: admin_academia exporta los usuarios de su propia academia a CSV.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = export_users(access, fmt='csv')
    assert rv.status_code == 200
    ctype = rv.headers.get('Content-Type', '')
    assert 'text/csv' in ctype


@pytest.mark.meta(title='Admin_academia - intento listar academia 303 debe fallar', desc='Intento de listar usuarios de academia 303 por admin_academia no vinculado debe 403')
def test_admin_academia_cannot_list_academia_303(client):
    print("\nPrueba: admin_academia intenta listar usuarios de academia_id=303 y debe recibir 403 o estar filtrado.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = list_users(access, params={'academia_id': 303})
    # Según la política de permisos esperamos 403 si intenta consultar otra academia
    assert rv.status_code == 403


@pytest.mark.meta(title='Admin_academia - intento export academia 303 debe fallar', desc='Intento de exportar usuarios de academia 303 por admin_academia no vinculado debe 403')
def test_admin_academia_cannot_export_academia_303_csv(client):
    print("\nPrueba: admin_academia intenta exportar usuarios de academia_id=303 a CSV y debe recibir 403.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = export_users(access, fmt='csv', params={'academia_id': 303})
    assert rv.status_code == 403


@pytest.mark.meta(title='Admin_academia - logout', desc='Logout del admin_academia')
def test_admin_academia_logout(client):
    print("\nPrueba: admin_academia realiza logout (revoca refresh token).")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = logout_with_refresh(refresh)
    assert rv.status_code in (200, 204)


# Test de paginación básica para Admin Academia
@pytest.mark.meta(title='Admin academia - paginación básica', desc='Comprobar paginación básica con page=1 y size=10')
def test_admin_academia_pagination_basic(client):
    print("\nPrueba: admin_academia verifica paginación básica (page=1, size=10).")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = list_users(access, params={'page': 1, 'size': 10})
    assert rv.status_code == 200
    resp = rv.json()
    assert len(resp['items']) <= 10
    assert resp['page'] == 1
    assert resp['size'] == 10
    assert isinstance(resp['has_more'], bool)


# Test de paginación avanzada para Admin Academia
@pytest.mark.meta(title='Admin academia - paginación avanzada', desc='Comprobar paginación avanzada con page=2 y size=5')
def test_admin_academia_pagination_advanced(client):
    print("\nPrueba: admin_academia verifica paginación avanzada (page=2, size=5).")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = list_users(access, params={'page': 2, 'size': 5})
    assert rv.status_code == 200
    resp = rv.json()
    assert len(resp['items']) <= 5
    assert resp['page'] == 2
    assert resp['size'] == 5
    assert isinstance(resp['has_more'], bool)


# Test de paginación sin permisos para Admin Academia
@pytest.mark.meta(title='Admin academia - paginación sin permisos', desc='Intentar paginar usuarios de academia_id=303 y recibir 403')
def test_admin_academia_pagination_no_permissions(client):
    print("\nPrueba: admin_academia intenta paginar usuarios de academia_id=303 y recibe 403.")
    access, refresh = login_and_get_tokens('admin_academia@academia.com', 'password_admin_academia')
    rv = list_users(access, params={'page': 1, 'size': 5, 'academia_id': 303})
    assert rv.status_code == 403
